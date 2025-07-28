import os
import re
import json
import tempfile
import shutil
import unicodedata
import traceback
import requests
from typing import Tuple, List, Dict
from pathlib import Path
from docx import Document
from collections import defaultdict
from urllib.parse import urljoin
import sys

sys.stdout.reconfigure(encoding='utf-8') if hasattr(sys.stdout, 'reconfigure') else None

class AnythingLLMProcessor:
    def __init__(self):
        pass

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "file_path": ("STRING", {"default": "", "file_input": True}),
                "api_url": ("STRING", {"default": "http://127.0.0.1:3001"}),
                "api_key": ("STRING", {"default": "your-api-key-here"}),
                "prompt_text": ("STRING", {
                    "multiline": True,
                    "default": "你是一位资深的知识库构建专家,专注于通过数据反推生成高质量问答对,你擅长从结构化或非结构化数据中提取核心信息,并据此构建适用于RAG（Retrieval-Augmented Generation）系统的知识库内容.请基于数据生成反推问题,问题类型需要包含事实性问题、推理性问题、对比性问题和总结性问题等多种类型,确保覆盖核心知识点和用户可能提出的真实查询场景.请基于以下内容生成问答对.格式为:\nQ: 问题\nA: 答案\n\n内容:\n"
                }),
                "chunk_limit": ("INT", {"default": 2250, "min": 100, "max": 10000}),
            },
            "optional": {
                "mode": (["query", "chat"], {"default": "query"}),
                "timeout": ("INT", {"default": 90, "min": 5, "max": 300}),
                "debug": ("BOOLEAN", {"default": True}),
                "workspace_slug": ("STRING", {"default": "default"}),
                "output_dir": ("STRING", {"default": "", "folder_input": True}),
            }
        }

    RETURN_TYPES = ("STRING",)
    RETURN_NAMES = ("status",)
    FUNCTION = "process_documents"
    CATEGORY = "RDC/问答对"

    def clean_path(self, path: str) -> str:
        """清理路径字符串，去除多余的引号和空格"""
        cleaned = path.strip().strip('"').strip("'")
        if os.name == 'nt':
            cleaned = cleaned.replace('\\\\', '\\')
        return os.path.normpath(cleaned)

    def normalize_filename(self, filename: str) -> str:
        """更安全的文件名规范化方法"""
        # 保留更多合法字符
        normalized = re.sub(r'[\\/:*?"<>|\x00-\x1F]', '_', filename)
        normalized = re.sub(r'\.{2,}', '.', normalized)  # 替换连续的点
        
        # 处理特殊前缀
        if re.match(r'^\d+[._]', normalized):
            normalized = f"chunk_{normalized}"
        
        # 移除首尾特殊字符
        normalized = normalized.strip('. _-')
        
        # 确保至少包含一个有效字符
        if not normalized:
            return "untitled"
        
        # 截断过长的文件名
        return normalized[:100]

    def get_safe_save_path(self, output_dir: str, base_name: str, index: int) -> str:
        """更健壮的文件保存路径生成"""
        # 处理空输出目录
        if not output_dir.strip():
            output_dir = os.path.join(os.getcwd(), "anythingllm_output")
        
        # 确保输出目录存在
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        
        # 生成唯一文件名
        base_name = self.normalize_filename(base_name)
        if not base_name:
            base_name = f"chunk_{index}"
        
        # 添加索引防止覆盖
        file_name = f"{base_name}_{index}.txt"
        save_path = os.path.join(output_dir, file_name)
        
        # 处理Windows特殊路径
        if os.name == 'nt':
            # 修复: 处理盘符根目录
            if re.match(r'^[a-zA-Z]:\\$', output_dir):
                save_path = output_dir + file_name
            # 处理短路径问题
            if len(save_path) > 260:
                # 使用临时目录缩短路径
                save_path = os.path.join(tempfile.gettempdir(), f"anythingllm_{index}.txt")
        
        # 确保路径长度在限制内
        if len(save_path) > 255:
            # 生成短哈希文件名
            import hashlib
            hash_name = hashlib.md5(base_name.encode('utf-8')).hexdigest()[:12]
            save_path = os.path.join(output_dir, f"{hash_name}_{index}.txt")
        
        return save_path

    def process_documents(self, file_path: str, api_url: str, api_key: str, 
                         prompt_text: str, chunk_limit: int, mode="query", 
                         timeout=90, debug=True, workspace_slug="default", output_dir=""):
        
        # 清理路径
        file_path = self.clean_path(file_path)
        output_dir = self.clean_path(output_dir) if output_dir else ""
        
        if not os.path.exists(file_path):
            return (f"错误: 文件不存在 {file_path}",)
        
        try:
            if debug:
                print(f"开始处理文档: {file_path}")
                print(f"使用工作区: {workspace_slug}")
                if output_dir:
                    print(f"指定输出目录: {output_dir}")
            
            # 验证API连接
            api_result = self.verify_api_connection(api_url, api_key, timeout, debug)
            if not api_result[0]:
                return (f"API连接失败: {api_result[1]}",)
            
            # 验证工作区是否存在
            if not self.verify_workspace(api_url, api_key, workspace_slug, timeout, debug):
                available_workspaces = self.get_available_workspaces(api_url, api_key, timeout, debug)
                return (f"错误: 工作区 '{workspace_slug}' 不存在. 可用工作区: {', '.join(available_workspaces)}",)
            
            # 解析文档结构
            ext = os.path.splitext(file_path)[1].lower()
            if debug:
                print(f"解析文档格式: {ext}")
            
            if ext == ".docx":
                document_structure = self.parse_docx(file_path, debug)
            elif ext in [".md", ".markdown"]:
                document_structure = self.parse_markdown(file_path, debug)
            else:
                return ("错误: 不支持的文件格式",)
            
            if not document_structure or not isinstance(document_structure, list):
                if debug:
                    print(f"文档解析失败: 类型={type(document_structure)}")
                return ("错误: 文档解析失败",)
            
            if debug:
                print(f"文档解析成功! 区块数量: {len(document_structure)}")
                # 打印前5个区块标题
                for i, section in enumerate(document_structure[:5]):
                    print(f"区块 {i+1} 标题: {section.get('title', '无标题')}")
            
            # 分块处理 - 使用新的分块逻辑
            chunks = self.chunk_document_with_hierarchy(document_structure, chunk_limit, debug)
            
            if debug:
                print(f"分块完成! 生成 {len(chunks)} 个区块")
                for i, chunk in enumerate(chunks):
                    content = "\n".join(chunk["content"])
                    title = chunk.get("title", f"区块_{i+1}")
                    print(f"区块 {i+1} ('{title}'): {len(content)} 字符")
            
            # 处理问答对
            saved_count = self.process_qa_chunks(
                chunks, api_url, api_key, workspace_slug, 
                prompt_text, mode, timeout, debug, output_dir)
            
            return (f"处理完成！保存了 {saved_count} 个问答文件",)
        except Exception as e:
            error_msg = f"处理错误: {str(e)}"
            if debug:
                print(f"详细错误信息:\n{traceback.format_exc()}")
            return (error_msg,)
    
    def verify_api_connection(self, api_url: str, api_key: str, timeout: int, debug: bool) -> Tuple[bool, str]:
        try:
            auth_url = urljoin(api_url, "/api/v1/auth")
            headers = {"Authorization": f"Bearer {api_key}"}
            
            response = requests.get(auth_url, headers=headers, timeout=timeout)
            
            if response.status_code == 200:
                result = response.json()
                if result.get("authenticated", False):
                    return (True, "验证成功")
                return (False, "API密钥无效或未授权")
            return (False, f"API返回错误状态码: {response.status_code}")
        except Exception as e:
            return (False, f"API连接错误: {str(e)}")

    def verify_workspace(self, api_url: str, api_key: str, workspace_slug: str, timeout: int, debug: bool) -> bool:
        try:
            workspaces_url = urljoin(api_url, "/api/v1/workspaces")
            headers = {"Authorization": f"Bearer {api_key}"}
            
            response = requests.get(workspaces_url, headers=headers, timeout=timeout)
            
            if response.status_code == 200:
                workspaces = response.json().get("workspaces", [])
                return any(ws['slug'] == workspace_slug for ws in workspaces)
            return False
        except Exception:
            return False
    
    def get_available_workspaces(self, api_url: str, api_key: str, timeout: int, debug: bool) -> List[str]:
        try:
            workspaces_url = urljoin(api_url, "/api/v1/workspaces")
            headers = {"Authorization": f"Bearer {api_key}"}
            
            response = requests.get(workspaces_url, headers=headers, timeout=timeout)
            
            if response.status_code == 200:
                workspaces = response.json().get("workspaces", [])
                return [ws['slug'] for ws in workspaces]
            return []
        except Exception:
            return []
    
    def parse_docx(self, file_path: str, debug: bool) -> List[Dict]:
        try:
            doc = Document(file_path)
            structure = []
            current_section = None
            current_level = 0
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                
                # 改进的标题检测逻辑
                level = 0
                if para.style.name.startswith('Heading'):
                    try:
                        level = int(para.style.name.split(' ')[1])
                    except:
                        # 如果无法提取级别，根据字体大小判断
                        if para.runs and para.runs[0].font.size:
                            if para.runs[0].font.size.pt > 20:
                                level = 1
                            elif para.runs[0].font.size.pt > 16:
                                level = 2
                            elif para.runs[0].font.size.pt > 14:
                                level = 3
                            else:
                                level = 4
                        else:
                            level = 1
                
                if level > 0:
                    # 创建新标题部分
                    current_section = {
                        "type": "heading",
                        "level": level,
                        "title": text,
                        "content": []
                    }
                    structure.append(current_section)
                    current_level = level
                else:
                    # 添加到当前部分
                    if current_section:
                        current_section["content"].append(text)
                    else:
                        # 创建默认部分
                        current_section = {
                            "type": "heading",
                            "level": 1,
                            "title": "Untitled",
                            "content": [text]
                        }
                        structure.append(current_section)
                        current_level = 1
            
            return structure
        except Exception as e:
            if debug:
                print(f"DOCX解析错误: {str(e)}")
                traceback.print_exc()
            return []

    def parse_markdown(self, file_path: str, debug: bool) -> List[Dict]:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
            
            # 改进的Markdown解析
            structure = []
            current_section = None
            
            # 按行处理
            lines = content.splitlines()
            i = 0
            
            while i < len(lines):
                line = lines[i].strip()
                if not line:
                    i += 1
                    continue
                
                # 检查标题行
                if line.startswith('#'):
                    # 提取标题级别
                    level = 0
                    while level < len(line) and line[level] == '#':
                        level += 1
                    
                    # 提取标题文本
                    title_text = line[level:].strip()
                    
                    # 确保有有效的标题
                    if not title_text:
                        title_text = f"标题_{len(structure)+1}"
                    
                    # 创建新标题部分
                    current_section = {
                        "type": "heading",
                        "level": level,
                        "title": title_text,
                        "content": []
                    }
                    structure.append(current_section)
                    i += 1
                    continue
                
                # 添加到当前部分
                if current_section:
                    current_section["content"].append(line)
                else:
                    # 创建默认部分（针对无标题的文档）
                    current_section = {
                        "type": "heading",
                        "level": 1,
                        "title": "Untitled",
                        "content": [line]
                    }
                    structure.append(current_section)
                
                i += 1
            
            return structure
        except Exception as e:
            if debug:
                print(f"Markdown解析错误: {str(e)}")
                traceback.print_exc()
            return []

    def chunk_document_with_hierarchy(self, structure: List[Dict], limit: int, debug: bool) -> List[Dict]:
        """根据标题层级分块文档，保留层级结构"""
        if not structure:
            return []
        
        chunks = []
        current_chunk = {"title": "初始区块", "content": [], "level": 1}
        current_size = 0
        parent_titles = {}  # 存储各级别标题
        
        for section in structure:
            level = section.get("level", 1)
            title = section.get("title", "未命名区块")
            content_lines = section.get("content", [])
            
            # 更新当前级别的标题
            parent_titles[level] = title
            
            # 构建层级标题路径
            title_path = []
            for lvl in sorted(parent_titles.keys()):
                if lvl <= level:
                    title_path.append(parent_titles[lvl])
            full_title = " > ".join(title_path)
            
            # 构建部分内容
            section_header = f"{'#' * level} {title}"
            section_content = [section_header] + content_lines
            section_text = "\n".join(section_content)
            section_size = len(section_text)
            
            # 检查是否超出限制
            if current_size + section_size <= limit:
                # 添加到当前块
                current_chunk["content"].extend(section_content)
                current_chunk["title"] = full_title
                current_chunk["level"] = level
                current_size += section_size
            else:
                # 保存当前块并创建新块
                if current_chunk["content"]:
                    chunks.append(current_chunk.copy())
                
                # 创建新块 - 继承父级标题
                current_chunk = {
                    "title": full_title,
                    "content": section_content,
                    "level": level
                }
                current_size = section_size
        
        # 添加最后一个块
        if current_chunk["content"]:
            chunks.append(current_chunk)
        
        return chunks

    def process_qa_chunks(self, chunks: List[Dict], api_url: str, api_key: str, 
                         workspace_slug: str, base_prompt: str, mode: str, 
                         timeout: int, debug: bool, output_dir: str) -> int:
        if not chunks:
            if debug:
                print("没有可处理的文档块")
            return 0
            
        saved_count = 0
        
        # 确定输出目录
        if not output_dir:
            output_dir = os.path.join(os.getcwd(), "anythingllm_output")
        
        os.makedirs(output_dir, exist_ok=True)
        
        if debug:
            print(f"输出目录: {output_dir}")
        
        # 构建API端点
        chat_url = urljoin(api_url, f"/api/v1/workspace/{workspace_slug}/chat")
        headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "Accept": "application/json"
        }
        
        if debug:
            print(f"API端点: {chat_url}")
            print(f"使用工作区: {workspace_slug}")

        for i, chunk in enumerate(chunks):
            content = "\n".join(chunk["content"])
            # 确保区块有标题
            title = chunk.get("title", f"区块_{i+1}")
            
            if debug:
                print(f"\n处理区块 #{i+1}: '{title}'")
                print(f"区块内容大小: {len(content)} 字符")
            
            # 生成安全文件名
            save_path = self.get_safe_save_path(output_dir, title, i+1)
            
            if debug:
                print(f"保存路径: {save_path}")
                
            # 构建完整提示
            full_prompt = f"{base_prompt}\n\n{content}"
                
            if debug:
                print(f"调用API生成QA对...")
                print(f"提示大小: {len(full_prompt)} 字符")
            
            # 调用API
            qa_text = self.call_anythingllm_api(chat_url, headers, full_prompt, mode, timeout, debug)
            
            if qa_text:
                if debug:
                    print(f"成功生成QA对! 大小: {len(qa_text)} 字符")
                    print(f"QA对预览: {qa_text[:200]}...")
                
                # 安全保存结果
                try:
                    with open(save_path, "w", encoding="utf-8") as f:
                        f.write(qa_text)
                    saved_count += 1
                    
                    if debug:
                        print(f"已保存到: {save_path}")
                except Exception as e:
                    if debug:
                        print(f"保存文件失败: {str(e)}")
            else:
                if debug:
                    print(f"警告: 区块 '{title}' 未能生成问答对")
        
        return saved_count

    def call_anythingllm_api(self, url: str, headers: dict, content: str, 
                            mode: str, timeout: int, debug: bool) -> str:
        try:
            # 构建请求体
            payload = {
                "message": content,
                "mode": mode,
                "stream": False
            }
            
            if debug:
                print(f"发送API请求 (模式: {mode})...")
                debug_content = content if len(content) < 500 else content[:500] + "..."
                print(f"内容预览: {debug_content}")
                print(f"请求体大小: {len(json.dumps(payload, ensure_ascii=False))} 字符")
            
            # 使用POST请求
            response = requests.post(
                url,
                headers=headers,
                json=payload,
                timeout=timeout
            )
            
            if debug:
                print(f"API响应状态: {response.status_code}")
                print(f"响应内容预览: {response.text[:500]}...")
            
            if response.status_code == 200:
                try:
                    result = response.json()
                    
                    # 修复: 根据实际API响应结构调整
                    # 新逻辑: 检查textResponse字段
                    if "textResponse" in result:
                        qa_text = result["textResponse"]
                        if qa_text:
                            return qa_text
                        elif debug:
                            print("API返回空textResponse")
                    # 兼容旧版本API结构
                    elif "response" in result:
                        qa_text = result["response"]
                        if qa_text:
                            return qa_text
                        elif debug:
                            print("API返回空response")
                    else:
                        if debug:
                            print("API返回无效结构，缺少textResponse/response字段")
                            print(f"完整响应: {result}")
                except json.JSONDecodeError:
                    if debug:
                        print(f"API返回无效JSON: {response.text[:300]}")
                    return ""
            else:
                if debug:
                    print(f"API错误详情: {response.text[:300]}")
                return ""
        except requests.exceptions.Timeout:
            if debug:
                print(f"API调用超时 (设置: {timeout}秒)")
            return ""
        except Exception as e:
            if debug:
                print(f"API调用失败: {str(e)}")
            return ""

    
# COMFYUI节点注册
NODE_CLASS_MAPPINGS = {
    "AnythingLLMProcessor": AnythingLLMProcessor
}

NODE_DISPLAY_NAME_MAPPINGS = {
    "AnythingLLMProcessor": "问答对处理"
}

CATEGORY_MAPPINGS = {
    "RDC/文档处理": "文档处理"
}

def get_custom_categories():
    return CATEGORY_MAPPINGS